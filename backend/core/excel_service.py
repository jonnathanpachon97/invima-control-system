from io import BytesIO

from openpyxl import Workbook

from openpyxl.styles import PatternFill, Font


def generate_records_excel(records):

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Registros"

    green_fill = PatternFill(
        start_color="C6EFCE",
        end_color="C6EFCE",
        fill_type="solid"
    )

    red_fill = PatternFill(
        start_color="FFC7CE",
        end_color="FFC7CE",
        fill_type="solid"
    )

    yellow_fill = PatternFill(
        start_color="FFF3CD",
        end_color="FFF3CD",
        fill_type="solid"
    )

    bold_font = Font(bold=True)

    headers = [
        "ID",
        "Fecha",
        "Formato",
        "Estado",
        "Observación",
        "Datos"
    ]

    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = bold_font

    for record in records:

        data_text = " | ".join([
            f"{key}: {value}"
            for key, value in record.data.items()
        ])

        sheet.append([
            record.id,

            record.created_at.strftime(
                "%Y-%m-%d %H:%M"
            ),

            record.template.name,

            record.status,

            record.observation or "",

            data_text
        ])

        status_cell = sheet.cell(
            row=sheet.max_row,
            column=4
        )

        if record.status == "aprobado":
            status_cell.fill = green_fill

        elif record.status == "rechazado":
            status_cell.fill = red_fill

        else:
            status_cell.fill = yellow_fill

    for column_cells in sheet.columns:

        length = max(
            len(str(cell.value))
            if cell.value else 0
            for cell in column_cells
        )

        adjusted_width = length + 5

        sheet.column_dimensions[
            column_cells[0].column_letter
        ].width = adjusted_width

    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer